"""Import teams, slogans, round scores, and source rankings from the TLC scorecard."""

import argparse
import asyncio
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.db.models import ClassTrack, EventEdition, Round, ScoreEntry, Team, TeamClass
from app.db.session import SessionFactory, dispose_engine


def read_rows(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    records: list[dict[str, object]] = []
    for class_code, sheet_name in (("IX", "Class IX"), ("X", "Class X"), ("XI", "Class XI")):
        sheet = workbook[sheet_name]
        for row in range(15, (sheet.max_row or 0) + 1):
            name = sheet.cell(row, 1).value
            slogan = sheet.cell(row, 2).value
            if not isinstance(name, str) or not name.strip():
                continue
            records.append({
                "class_code": class_code,
                "name": name.strip(),
                "slogan": str(slogan).strip() if slogan is not None else None,
                "round0": sheet.cell(row, 4).value,
                "round1": sheet.cell(row, 6).value,
                "total": sheet.cell(row, 15).value,
                "out_of": sheet.cell(row, 16).value,
                "percent": sheet.cell(row, 17).value,
                "rank": sheet.cell(row, 18).value,
            })
    return records


async def import_scorecard(path: Path) -> None:
    records = read_rows(path)
    if not records:
        raise RuntimeError("No team rows were found in the scorecard workbook.")

    async with SessionFactory() as session:
        edition = await session.scalar(select(EventEdition).where(EventEdition.event_year == 2026))
        if edition is None:
            raise RuntimeError("Event edition 2026 is missing. Run seed_demo_data.py first.")
        tracks = {item.code: item for item in (await session.scalars(select(ClassTrack).where(ClassTrack.event_edition_id == edition.id))).all()}
        rounds: dict[int, Round] = {}
        for number, name, round_type in ((0, "Team Identity", "OTHER"), (1, "Presentation", "PRESENTATION")):
            item = await session.scalar(select(Round).where(Round.event_edition_id == edition.id, Round.round_number == number))
            if item is None:
                item = Round(event_edition_id=edition.id, round_number=number, name=name, round_type=round_type, status="ACTIVE")
                session.add(item)
                await session.flush()
            rounds[number] = item

        await session.execute(delete(ScoreEntry).where(ScoreEntry.note.like("Imported from TECH_EVENT_SCORECARD:%")))
        imported_teams: set[str] = set()
        for record in records:
            track = tracks[record["class_code"]]
            team = await session.scalar(select(Team).where(Team.event_edition_id == edition.id, Team.name == record["name"]))
            if team is None:
                team = Team(event_edition_id=edition.id, name=record["name"], slogan=record["slogan"], team_type="LEARNING", is_public=True)
                session.add(team)
                await session.flush()
            else:
                team.slogan = record["slogan"]
                team.is_public = True
            if await session.scalar(select(TeamClass).where(TeamClass.team_id == team.id, TeamClass.class_track_id == track.id)) is None:
                session.add(TeamClass(team_id=team.id, class_track_id=track.id))
            imported_teams.add(team.name)
            source_note = f"Imported from TECH_EVENT_SCORECARD: {record['class_code']} rank {record['rank']} total {record['total']}/{record['out_of']} ({record['percent']})"
            for number, value in ((0, record["round0"]), (1, record["round1"])):
                if isinstance(value, (int, float)):
                    session.add(ScoreEntry(team_id=team.id, round_id=rounds[number].id, points_delta=value, reason="ADJUSTMENT", note=source_note))

        demo = await session.scalar(select(Team).where(Team.event_edition_id == edition.id, Team.name == "Demo Team"))
        if demo is not None and demo.name not in imported_teams:
            demo.is_public = False
        await session.commit()
        print(f"Imported {len(records)} teams and score entries for {edition.name}.")


async def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    try:
        await import_scorecard(args.workbook)
    finally:
        await dispose_engine()


if __name__ == "__main__":
    asyncio.run(main())
